#!/usr/bin/env python3

from openpyxl import load_workbook
import pandas as pd
import yaml
import sys

INPUT_FILE = sys.argv[1] if len(sys.argv) > 1 else "Overview.xlsx"
PROJECT_OUTPUT = sys.argv[2] if len(sys.argv) > 2 else "Overview.yaml"
COLOURS_OUTPUT = sys.argv[3] if len(sys.argv) > 3 else "Colours.yaml"


# ======================================================
# Helpers
# ======================================================

def clean_date(value):

    if pd.isna(value):
        return None

    return pd.to_datetime(value).date()


def clean_text(value):

    if pd.isna(value):
        return ""

    return str(value).strip()


def clean_int(value, default=0):

    if pd.isna(value):
        return default

    try:
        return int(value)
    except Exception:
        return default


def split_risks(value):

    if pd.isna(value):
        return []

    value = str(value).strip()

    if value == "":
        return []

    return [
        item.strip()
        for item in value.split(";")
        if item.strip()
    ]


def excel_fill_to_hex(cell):

    colour = cell.fill.fgColor

    # standard RGB
    if colour.type == "rgb":

        rgb = colour.rgb

        if rgb:

            if len(rgb) == 8:
                rgb = rgb[2:]

            return f"#{rgb}"

    # indexed colours
    elif colour.type == "indexed":
        print(
            f"Error xlsx colour indexed {colour_cell}"
        )
        return None

    # theme colours
    elif colour.type == "theme":
        print(
            f"Error xlsx colour theme {colour_cell}"
        )
        return None

    return None


def get_colour(cell, fallback):

    # Prefer explicit hex value
    if cell.value is not None:

        value = str(cell.value).strip()

        if value.startswith("#"):
            return value

    # Otherwise try cell fill
    fill = excel_fill_to_hex(cell)

    if fill:
        return fill

    # Fallback if neither exists
    return fallback


# ======================================================
# Read workbook
# ======================================================

wb = load_workbook(INPUT_FILE)
partners_ws = wb["Partners"]

wps_df = pd.read_excel(
    INPUT_FILE,
    sheet_name="WPs"
)

tasks_df = pd.read_excel(
    INPUT_FILE,
    sheet_name="Tasks"
)

deps_df = pd.read_excel(
    INPUT_FILE,
    sheet_name="Dependencies"
)

partners_df = pd.read_excel(
    INPUT_FILE,
    sheet_name="Partners"
)

statuses_df = pd.read_excel(
    INPUT_FILE,
    sheet_name="Statuses"
)

# remove accidental whitespace from columns

for df in [
    wps_df,
    tasks_df,
    deps_df,
    partners_df,
    statuses_df
]:

    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
    )

# ======================================================
# Build Overview.yaml
# ======================================================

overview = {
    "work_packages": {},
    "dependencies": []
}

# ------------------------
# Work Packages
# ------------------------

for _, row in wps_df.iterrows():

    wp = clean_text(row["WP"])

    overview["work_packages"][wp] = {

        "title": clean_text(row["Title"]),

        "lead": clean_text(row["Lead"]),

        "deliverables": {},

        "milestones": {}
    }

# ------------------------
# Tasks
# ------------------------

for _, row in tasks_df.iterrows():

    wp = clean_text(row["WP"])
    item_id = clean_text(row["ID"])

    task = {

        "status": clean_text(row["Status"]),

        "progress": clean_int(row.get("Progress", 0)),

        "title": clean_text(row["Title"]),

        "lead": clean_text(row["Lead"]),

        "start": clean_date(row["Start"]),

        "due": clean_date(row["Due"]),

        "notes": clean_text(
            row.get("Notes", "")
        ),

        "risks": split_risks(
            row.get("Risks", "")
        )
    }

    if item_id.startswith("M"):

        overview["work_packages"][wp][
            "milestones"
        ][item_id] = task

    else:

        overview["work_packages"][wp][
            "deliverables"
        ][item_id] = task

# ------------------------
# Dependencies
# ------------------------

for _, row in deps_df.iterrows():

    from_id = clean_text(
        row["From"]
    )

    to_id = clean_text(
        row["To"]
    )

    if from_id and to_id:

        overview["dependencies"].append(
            {
                "from": from_id,
                "to": to_id
            }
        )

# ======================================================
# Build Colours.yaml
# ======================================================

colours = {

    "partners": {},

    "statuses": {},

    "milestone": "#111827",

    "fallback": {
        "node": "#ffffff",
        "cluster": "#eeeeee"
    }
}

# ------------------------
# Partners
# ------------------------

headers = {}

for col_num, cell in enumerate(partners_ws[1], start=1):
    headers[cell.value] = col_num

for row_num in range(2, partners_ws.max_row + 1):

    partner = clean_text(
        partners_ws.cell(
            row=row_num,
            column=headers["Partner"]
        ).value
    )

    node_cell = partners_ws.cell(
        row=row_num,
        column=headers["NodeColour"]
    )

    cluster_cell = partners_ws.cell(
        row=row_num,
        column=headers["ClusterColour"]
    )

    colours["partners"][partner] = {

        "node": get_colour(
            node_cell,
            colours["fallback"]["node"]
        ),

        "cluster": get_colour(
            cluster_cell,
            colours["fallback"]["cluster"]
        ),

        "contact": clean_text(
            partners_ws.cell(
                row=row_num,
                column=headers.get("Contact")
            ).value
        ),

        "email": clean_text(
            partners_ws.cell(
                row=row_num,
                column=headers.get("Email")
            ).value
        ),

        "phone": clean_text(
            partners_ws.cell(
                row=row_num,
                column=headers.get("Phone")
            ).value
        )
    }

# ------------------------
# Statuses
# ------------------------

statuses_ws = wb["Statuses"]

status_headers = {}

for col_num, cell in enumerate(
    statuses_ws[1],
    start=1
):
    status_headers[cell.value] = col_num

for row_num in range(
    2,
    statuses_ws.max_row + 1
):

    status = clean_text(
        statuses_ws.cell(
            row=row_num,
            column=status_headers["Status"]
        ).value
    )

    colour_cell = statuses_ws.cell(
        row=row_num,
        column=status_headers["BorderColour"]
    )

    colours["statuses"][status] = get_colour(
        colour_cell,
        "#6b7280"
    )

# ======================================================
# Write files
# ======================================================

with open(
    PROJECT_OUTPUT,
    "w",
    encoding="utf-8"
) as f:

    yaml.dump(
        overview,
        f,
        sort_keys=False,
        default_flow_style=False,
        allow_unicode=True
    )

with open(
    COLOURS_OUTPUT,
    "w",
    encoding="utf-8"
) as f:

    yaml.dump(
        colours,
        f,
        sort_keys=False,
        default_flow_style=False,
        allow_unicode=True
    )

print(
    f"Created {PROJECT_OUTPUT}"
)

print(
    f"Created {COLOURS_OUTPUT}"
)